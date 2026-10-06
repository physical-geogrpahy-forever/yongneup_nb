#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PB4 CHELSA21K full-output exporter.

This is a post-processing utility only. It does not modify PB4 science calculations.
It reads the completed outputs_CHELSA21K directory and writes:
  - combined 211-time-step CSVs with every summary_timeseries column
  - a major-variable Excel workbook
  - validation CSV
  - output-file manifest with SHA-256
  - time-series PNG figures
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt


MAJOR_COLUMNS = [
    "mode", "ka_bp",
    "mean_elevation_m", "sd_elevation_m", "min_elevation_m", "max_elevation_m", "relief_m",
    "mean_soil_depth_m", "sd_soil_depth_m", "min_soil_depth_m", "max_soil_depth_m",
    "mean_slope", "sd_slope", "min_slope", "max_slope",
    "mean_npp", "sd_npp", "min_npp", "max_npp",
    "mean_eemt", "sd_eemt", "min_eemt", "max_eemt",
    "reich_lai_sapwood_agb_mean_kg_m2",
    "reich_lai_sapwood_agb_median_kg_m2",
    "reich_lai_sapwood_agb_p05_kg_m2",
    "reich_lai_sapwood_agb_p95_kg_m2",
    "reich_lai_sapwood_agb_max_kg_m2",
    "dominant_vegetation_code", "vegetation_class_richness",
    "vegetation_class_counts",
    "geomorph_accepted_substeps_to_next",
    "geomorph_rejected_substeps_to_next",
    "geomorph_rejected_fluvial_trials_to_next",
    "geomorph_min_dt_kyr_to_next",
    "geomorph_max_dt_kyr_to_next",
    "geomorph_fluvial_tolerance_m",
    "geomorph_max_accepted_fluvial_change_m_to_next",
]

VARIABLE_DICTIONARY = [
    ("ka_bp", "모델 시점", "ka BP", "21.0 ka BP부터 0.0 ka BP까지 0.1 kyr 간격"),
    ("mean_elevation_m", "유역 평균 지표고도", "m", "유효 격자의 평균 z"),
    ("sd_elevation_m", "지표고도 표준편차", "m", "유효 격자 간 공간 표준편차"),
    ("min_elevation_m", "최저 지표고도", "m", "유효 격자 최소값"),
    ("max_elevation_m", "최고 지표고도", "m", "유효 격자 최대값"),
    ("relief_m", "기복", "m", "max elevation - min elevation"),
    ("mean_soil_depth_m", "평균 토심", "m", "유효 격자의 평균 H"),
    ("sd_soil_depth_m", "토심 표준편차", "m", "유효 격자 간 공간 표준편차"),
    ("min_soil_depth_m", "최소 토심", "m", "유효 격자 최소값"),
    ("max_soil_depth_m", "최대 토심", "m", "유효 격자 최대값"),
    ("mean_slope", "평균 경사", "gradient", "PB4 raster slope"),
    ("mean_npp", "평균 NPP", "gC m-2 yr-1", "BIOME4 NPP"),
    ("mean_eemt", "평균 EEMT", "MJ m-2 yr-1", "Pelletier coupling에 사용되는 EEMT"),
    ("reich_lai_sapwood_agb_mean_kg_m2", "평균 AGB*", "kg dry biomass m-2", "Reich leaf + Haxeltine/Prentice living sapwood proxy"),
    ("dominant_vegetation_code", "우점 식생 코드", "code", "검증용 reduced vegetation class"),
    ("vegetation_class_counts", "식생 분류별 격자수", "count", "시점별 class:count 문자열"),
    ("geomorph_accepted_substeps_to_next", "다음 100년 구간 accepted substeps", "count", "adaptive geomorphic integration"),
    ("geomorph_rejected_substeps_to_next", "다음 100년 구간 rejected substeps", "count", "adaptive geomorphic integration"),
    ("geomorph_rejected_fluvial_trials_to_next", "fluvial rejected trials", "count", "adaptive fluvial stability diagnostic"),
    ("geomorph_min_dt_kyr_to_next", "최소 지형 substep", "kyr", "다음 100년 구간"),
    ("geomorph_max_dt_kyr_to_next", "최대 지형 substep", "kyr", "다음 100년 구간"),
    ("geomorph_fluvial_tolerance_m", "fluvial 변화 허용치", "m", "수치 안정성 설정"),
    ("geomorph_max_accepted_fluvial_change_m_to_next", "최대 accepted fluvial change", "m", "다음 100년 구간"),
]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_summary(outputs: Path, mode: str) -> pd.DataFrame:
    p = outputs / f"model_{mode}" / "summary_timeseries.csv"
    if not p.exists():
        raise FileNotFoundError(p)
    df = pd.read_csv(p, encoding="utf-8-sig")
    if "ka_bp" not in df.columns:
        raise RuntimeError(f"ka_bp missing: {p}")
    df = df.copy()
    if "mode" in df.columns:
        df["mode"] = mode
    else:
        df.insert(0, "mode", mode)
    return df


def verify_211(df: pd.DataFrame, mode: str) -> None:
    vals = np.sort(df["ka_bp"].astype(float).unique())
    if len(df) != 211 or len(vals) != 211:
        raise RuntimeError(f"{mode}: expected 211 rows/times, got rows={len(df)}, unique={len(vals)}")
    if not np.isclose(vals[0], 0.0) or not np.isclose(vals[-1], 21.0):
        raise RuntimeError(f"{mode}: time range is {vals[0]}..{vals[-1]}, expected 0..21 ka BP")
    diffs = np.diff(vals)
    if not np.allclose(diffs, 0.1, atol=1e-9, rtol=0):
        raise RuntimeError(f"{mode}: time grid is not uniform 0.1 kyr")


def collect_validation(outputs: Path) -> pd.DataFrame:
    frames = []
    for mode in ("static", "dynamic"):
        p = outputs / f"model_{mode}" / "validation_per_output_time.csv"
        if p.exists():
            d = pd.read_csv(p, encoding="utf-8-sig")
            if "mode" not in d.columns:
                d.insert(0, "mode", mode)
            frames.append(d)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def collect_climate(root: Path) -> pd.DataFrame:
    candidates = [
        root / "YONGNEUP_CHELSA_TRACE21k_ENVICLOUD_RAW_WIDE.csv",
        root / "data" / "YONGNEUP_CHELSA_TRACE21k_ENVICLOUD_RAW_WIDE.csv",
    ]
    for p in candidates:
        if p.exists():
            return pd.read_csv(p, encoding="utf-8-sig")
    return pd.DataFrame()


def collect_config(outputs: Path) -> pd.DataFrame:
    p = outputs / "run_config_CHELSA21K.json"
    if not p.exists():
        return pd.DataFrame()
    obj = json.loads(p.read_text(encoding="utf-8"))
    rows = []
    def walk(prefix, v):
        if isinstance(v, dict):
            for k, x in v.items():
                walk(f"{prefix}.{k}" if prefix else str(k), x)
        elif isinstance(v, list):
            rows.append((prefix, json.dumps(v, ensure_ascii=False)))
        else:
            rows.append((prefix, v))
    walk("", obj)
    return pd.DataFrame(rows, columns=["parameter", "value"])


def output_manifest(outputs: Path, export_dir: Path) -> pd.DataFrame:
    rows = []
    for p in sorted(outputs.rglob("*")):
        if not p.is_file() or export_dir in p.parents:
            continue
        try:
            digest = sha256_file(p)
        except Exception:
            digest = ""
        rows.append({
            "relative_path": p.relative_to(outputs).as_posix(),
            "suffix": p.suffix.lower(),
            "size_bytes": p.stat().st_size,
            "sha256": digest,
        })
    return pd.DataFrame(rows)


def make_plots(all_ts: pd.DataFrame, export_dir: Path) -> None:
    specs = [
        ("mean_elevation_m", "Mean elevation", "m", "mean_elevation_211times.png"),
        ("mean_soil_depth_m", "Mean soil depth", "m", "mean_soil_depth_211times.png"),
        ("mean_npp", "Mean NPP", "gC m-2 yr-1", "mean_npp_211times.png"),
        ("mean_eemt", "Mean EEMT", "MJ m-2 yr-1", "mean_eemt_211times.png"),
        ("reich_lai_sapwood_agb_mean_kg_m2", "Mean AGB*", "kg m-2", "mean_agb_star_211times.png"),
    ]
    for col, title, ylabel, name in specs:
        if col not in all_ts.columns:
            continue
        plt.figure(figsize=(9, 5.2))
        for mode in ("static", "dynamic"):
            d = all_ts[all_ts["mode"] == mode].sort_values("ka_bp", ascending=False)
            if len(d):
                plt.plot(d["ka_bp"], d[col], label=mode)
        plt.gca().invert_xaxis()
        plt.xlabel("ka BP")
        plt.ylabel(ylabel)
        plt.title(f"{title} through 21 ka")
        plt.legend()
        plt.tight_layout()
        plt.savefig(export_dir / name, dpi=220)
        plt.close()


def style_excel(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="D9EAF7")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            for row in range(1, min(ws.max_row, 80) + 1):
                v = ws.cell(row, col_idx).value
                max_len = max(max_len, len(str(v)) if v is not None else 0)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 34)

    if "Major_variables" in wb.sheetnames:
        ws = wb["Major_variables"]
        # create chart sheet with dynamic series from Major_variables
        if "Charts" not in wb.sheetnames:
            ch = wb.create_sheet("Charts")
        else:
            ch = wb["Charts"]
        ch["A1"] = "PB4 주요 변수 211시점 그래프"
        ch["A1"].font = Font(bold=True, size=14)
        # Charts use the dedicated dynamic sheet for clean x/y series.
        if "Dynamic_all" in wb.sheetnames:
            src = wb["Dynamic_all"]
            headers = {src.cell(1, c).value: c for c in range(1, src.max_column + 1)}
            xcol = headers.get("ka_bp")
            chart_specs = [
                ("mean_elevation_m", "평균 고도", "m", "A3", "H20"),
                ("mean_soil_depth_m", "평균 토심", "m", "J3", "Q20"),
                ("mean_npp", "평균 NPP", "gC m-2 yr-1", "A22", "H39"),
                ("mean_eemt", "평균 EEMT", "MJ m-2 yr-1", "J22", "Q39"),
                ("reich_lai_sapwood_agb_mean_kg_m2", "평균 AGB*", "kg m-2", "A41", "H58"),
            ]
            if xcol:
                for field, title, ytitle, start, end in chart_specs:
                    ycol = headers.get(field)
                    if not ycol:
                        continue
                    chart = LineChart()
                    chart.title = title
                    chart.y_axis.title = ytitle
                    chart.x_axis.title = "ka BP"
                    data = Reference(src, min_col=ycol, min_row=1, max_row=src.max_row)
                    cats = Reference(src, min_col=xcol, min_row=2, max_row=src.max_row)
                    chart.add_data(data, titles_from_data=True)
                    chart.set_categories(cats)
                    chart.height = 8
                    chart.width = 14
                    ch.add_chart(chart, start)

    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".", help="PB4 package root")
    ap.add_argument("--outputs", default="outputs_CHELSA21K")
    args = ap.parse_args()

    root = Path(args.root).resolve()
    outputs = (root / args.outputs).resolve()
    export_dir = outputs / "FULL_EXPORT"
    export_dir.mkdir(parents=True, exist_ok=True)

    static = read_summary(outputs, "static")
    dynamic = read_summary(outputs, "dynamic")
    verify_211(static, "static")
    verify_211(dynamic, "dynamic")

    # Full summary: every recorded column, every one of 211 times for each mode.
    all_cols = list(dict.fromkeys(list(static.columns) + list(dynamic.columns)))
    static_all = static.reindex(columns=all_cols)
    dynamic_all = dynamic.reindex(columns=all_cols)
    all_ts = pd.concat([static_all, dynamic_all], ignore_index=True)
    all_ts.to_csv(export_dir / "PB4_ALL_TIMESTEPS_ALL_VARIABLES.csv", index=False, encoding="utf-8-sig")
    static_all.to_csv(export_dir / "PB4_STATIC_211_TIMESTEPS_ALL_VARIABLES.csv", index=False, encoding="utf-8-sig")
    dynamic_all.to_csv(export_dir / "PB4_DYNAMIC_211_TIMESTEPS_ALL_VARIABLES.csv", index=False, encoding="utf-8-sig")

    major_cols = [c for c in MAJOR_COLUMNS if c in all_ts.columns]
    major = all_ts[major_cols].copy()
    major.to_csv(export_dir / "PB4_MAJOR_VARIABLES_211_TIMESTEPS.csv", index=False, encoding="utf-8-sig")

    validation = collect_validation(outputs)
    if not validation.empty:
        validation.to_csv(export_dir / "PB4_VALIDATION_ALL_TIMES.csv", index=False, encoding="utf-8-sig")

    climate = collect_climate(root)
    config = collect_config(outputs)
    manifest = output_manifest(outputs, export_dir)
    manifest.to_csv(export_dir / "PB4_OUTPUT_FILE_MANIFEST.csv", index=False, encoding="utf-8-sig")

    var_dict = pd.DataFrame(VARIABLE_DICTIONARY, columns=["variable", "설명", "단위", "비고"])

    xlsx = export_dir / "PB4_MAJOR_VARIABLES_211_TIMESTEPS.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        major.to_excel(writer, sheet_name="Major_variables", index=False)
        dynamic_all.to_excel(writer, sheet_name="Dynamic_all", index=False)
        static_all.to_excel(writer, sheet_name="Static_all", index=False)
        var_dict.to_excel(writer, sheet_name="Variable_dictionary", index=False)
        if not validation.empty:
            validation.to_excel(writer, sheet_name="Validation", index=False)
        if not climate.empty:
            # Excel row limit guard
            climate.iloc[:1048575].to_excel(writer, sheet_name="Climate_input", index=False)
        if not config.empty:
            config.to_excel(writer, sheet_name="Run_config", index=False)
        manifest.to_excel(writer, sheet_name="Output_manifest", index=False)

    style_excel(xlsx)
    make_plots(all_ts, export_dir)

    readme = export_dir / "README_FULL_EXPORT_KO.txt"
    readme.write_text(
        "PB4 CHELSA21K FULL EXPORT\n"
        "==========================\n"
        "이 폴더는 PB4 계산 이후 생성되는 후처리 산출물입니다.\n"
        "과학 계산 코드는 변경하지 않습니다.\n\n"
        "핵심 파일\n"
        "- PB4_MAJOR_VARIABLES_211_TIMESTEPS.xlsx : 주요 변수 Excel, static/dynamic 전체 211시점\n"
        "- PB4_ALL_TIMESTEPS_ALL_VARIABLES.csv : summary_timeseries의 모든 기록 열, 총 422행\n"
        "- PB4_DYNAMIC_211_TIMESTEPS_ALL_VARIABLES.csv : dynamic 211시점 전체 열\n"
        "- PB4_STATIC_211_TIMESTEPS_ALL_VARIABLES.csv : static 211시점 전체 열\n"
        "- PB4_OUTPUT_FILE_MANIFEST.csv : outputs_CHELSA21K 내 모든 원산출물 파일 목록과 SHA-256\n"
        "- *_211times.png : 주요 변수 전체시계열 그래프\n\n"
        "주의: 격자별 GeoTIFF/NPZ/CSV 등 원산출물은 outputs_CHELSA21K 아래 원래 위치에 그대로 보존됩니다.\n",
        encoding="utf-8"
    )

    print(f"PASS: static rows={len(static)}, dynamic rows={len(dynamic)}")
    print(f"FULL_EXPORT={export_dir}")
    print(f"EXCEL={xlsx}")
    print(f"ALL_COLUMNS={len(all_cols)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
